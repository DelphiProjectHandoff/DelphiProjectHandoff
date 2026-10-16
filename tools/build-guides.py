"""Build DPH manuals from maintained prose and exact source references."""
from pathlib import Path
from datetime import date
import hashlib,json,html,re,textwrap
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
CHANGED='16 October 2026'
VERSION=52
LAYER_PURPOSE={
'Project.Source':'Owns original bytes, decoding, character to byte ranges and verified acquisition. Preserve byte identity and ranges when changing encodings.',
'Project.Inventory':'Discovers the selected project and permitted dependencies under its configuration. Maintains source availability, defines and path boundaries.',
'Project.Analysis':'Coordinates syntax scans, incremental reuse, context identity and generation staging. Owns acquired inventory and parse-result lifetimes.',
'Parse.Model':'Defines tokens, ranges, diagnostics and intermediate syntax records. Schema/version and acquisition bounds are contracts for every downstream layer.',
'Parse.Lexer':'Converts acquired source text into ranged tokens. Malformed input and token limits must remain qualified rather than silently normalized.',
'Parse.Preprocessor':'Evaluates supported directives and three-valued conditional activity, including allowed include resolution. Unknown activity must survive projection.',
'Parse.Pascal':'Builds supported Pascal declarations and statement IR. Inspect parent indexes and source ranges when extending syntax coverage.',
'Parse.Forms':'Reads supported Delphi form definitions into component and property structure. Captured designs remain distinct from actual screenshots.',
'Parse.Resources':'Reads supported resource definitions and preserves qualified evidence for unsupported material. It does not establish runtime resource behavior.',
'Analysis.Model':'Owns semantic symbols, relations and control-flow data plus provenance. Unresolved relations and bounded traversal remain explicit.',
'Analysis.Symbols':'Builds and binds supported semantic references, including validated routine-summary reuse. Cache changes require full/incremental agreement.',
'Analysis.Flow':'Constructs statement control-flow topology with explicit continuations, exceptions, returns, breaks and finally context. Preserve branch semantics.',
'Analysis.Async':'Adds qualified evidence for supported asynchronous source patterns. Static patterns are not runtime scheduling observations.',
'Analysis.SQL':'Adds supported SQL/data-use evidence from source. It is not permission to connect to or execute statements against target data.',
'Analysis.Persistence':'Persists supported semantic facts and their references into generation-scoped storage. Keep fact transactions and publication separate.',
'Analysis.Quality':'Projects metrics and investigation candidates from captured facts. Candidate findings do not prove unused code or safe removal.',
'Analysis.Explanations':'Builds learner topics and object explanations from source facts and relevant attributed knowledge. Keep meaning first and deeper code optional.',
'Handoff.Dossier':'Renders the selected capture as a four-part, source-linked maintenance handoff. Distinguish extracted facts, author statements, reading guidance and unresolved questions.',
'Learning.Discovery':'Finds project-local documents, logs, saved screens, startup forms and referenced public websites. Retain source paths, bounded scanning, original-language text and explicit extraction status.',
'Storage.Types':'Defines persistent status and manifest types shared by analysis and storage. Changes require serialization and lifecycle compatibility checks.',
'Storage.Database':'Owns the SQLite writer, migrations, generations, knowledge revisions, artifacts and indexes. Thread ownership and transactional boundaries are mandatory.',
'Storage.Backup':'Creates and restores consistent database plus managed-artifact copies with hash checks. Fresh destinations and outside-source paths are required.',
'Workspace.Worker':'Serializes background commands and returns owned response objects. Preserve queue bounds, COM lifecycle and response ownership transfers.',
'Workspace.Catalog':'Stores local database navigation and revision-specific learner progress. Atomic replacement protects the existing preference file.',
'Workspace.Model':'Owns the captured workspace view and typed teaching-record validators/canonical renderers. Keep structured metadata and readable body consistent.',
'Workspace.Screen':'Projects source-based form geometry for inspection. Results teach captured layout and do not claim an observed runtime screenshot.',
'Workspace.Graphs':'Projects bounded graph views for investigation. Preserve source identity, edge qualification and selection routes.',
'Workspace.Comparison':'Compares saved captured evidence between generations. It is not reconstruction of unavailable historical source text.',
'Workspace.ScreenHistory':'Compares supplied/captured screen material with generation-qualified identity. Keep actual images and reconstructed forms distinct.',
'Workspace.ActionPurposes':'Projects attributed action-purpose records with supported source links. Business intent must remain supplied knowledge.',
'Workspace.ComponentDependencies':'Projects component dependency evidence for focused inspection. A source dependency is not proof of deployed runtime availability.',
'Workspace.DatabaseConnections':'Projects source-derived database connection declarations. Never interpret a declaration as successful connectivity.',
'Workspace.DatabaseSchema':'Projects supported schema definitions and supplied descriptions. It does not introspect a live target database automatically.',
'Workspace.Commands':'Projects commands, actions and handler connections. Keep command declarations separate from observed invocation.',
'Workspace.FormInheritance':'Projects supported saved-form inheritance. Incomplete ancestor material requires explicit qualification.',
'Workspace.DynamicControls':'Projects supported dynamic-control patterns. Source evidence cannot guarantee all runtime-created controls are reconstructed.',
'Workspace.Keyboard':'Projects keyboard/shortcut evidence and verified source navigation. Preserve detached-source limits and original ranges.',
'Workspace.LayoutRules':'Explains captured alignment, anchors and spacing. Educational resize projections are not observed runtime geometry.',
'Knowledge.ProjectProfile':'Validates and renders an attributed application profile. Purpose, support, safe practice and unknowns support orientation.',
'Knowledge.Glossary':'Validates business terms and aliases and produces canonical text. Normalized naming rules prevent conflicting current vocabulary.',
'Knowledge.ActionPurpose':'Validates stable action-purpose identity and twelve business explanation fields. Source targets and human intent remain separate.',
'Knowledge.ComponentDependency':'Validates attributed dependency descriptions and canonical presentation. Names do not prove deployment or licensing status.',
'Publish.Model':'Builds the scoped canonical document, audiences, provenance and redacted knowledge. This is the privacy boundary shared by output formats.',
'Publish.Export':'Writes fresh offline editions and validates checksums, provenance, ranges and indexes. Learner and technical artifacts share the scoped model.',
'Publish.OfflinePreview':'Serves the local preview and bundled German translation assets over loopback. It must not expose the selected project to an external network interface.',
'Publish.BrowserPolicy':'Constrains local embedded navigation and internal routes. Rendering success does not replace URL containment tests.',
'Publish.Translation':'Adapts pinned translation capabilities and revision/locale cache state. Provider configuration and named review remain separate.',
'UI.Main':'Binds serialized FireMonkey controls to worker requests and owned models. Preserve pending drafts, busy state, selection identity and close guards.',
'UI.ProjectProfile':'Collects application-profile fields in an IDE-editable dialog. Applying returns a draft; general knowledge saving persists it.',
'UI.Glossary':'Collects a business term and its supporting meaning/aliases. Validate before applying the draft to saved knowledge.',
'UI.ActionPurpose':'Collects an action explanation in a serialized dialog. Keep user purpose and supporting captured target clearly identified.',
'UI.ComponentDependency':'Collects attributed component/dependency descriptions. It is a knowledge editor, not an installer or dependency resolver.',
'UI.DatabaseDictionary':'Collects authored dictionary descriptions. This dialog does not become a live database administration engine.'}

def source_blocks(name):
 lines=(ROOT/'docs/guides/source'/name).read_text(encoding='utf-8-sig').splitlines();blocks=[];i=0
 while i<len(lines):
  line=lines[i].strip();i+=1
  if not line:continue
  if line in ['@steps','@code']:
   values=[]
   while i<len(lines) and lines[i].strip()!='@end':values.append(lines[i]);i+=1
   i+=1;blocks.append((line[1:],values));continue
  if line.startswith('@'):
   key,_,value=line[1:].partition(' ');blocks.append((key,value))
  else:blocks.append(('p',line))
 return blocks

def units_blocks():
 blocks=[]
 for i,path in enumerate(sorted(ROOT.glob('DPH.*.pas')),1):
  text=path.read_text(encoding='utf-8-sig');interface=text.split('\nimplementation',1)[0]
  name=path.stem.removeprefix('DPH.');assert name in LAYER_PURPOSE,name
  blocks.extend([('section',f'25.{i+1}. {name}'),('p','Source file: '+path.name+'. '+LAYER_PURPOSE[name])])
  types=re.findall(r'^\s*(TDPH\w+)\s*=\s*(?:class|record|\(|reference)',interface,re.M)
  if types:blocks.append(('p','Principal public types: '+', '.join(dict.fromkeys(types))+'.'))
  # Capture complete public/interface declarations across line wraps, without implementation bodies.
  # Parameter semicolons occur inside signatures; use a parenthesis-aware scan for complete declarations.
  declarations=[];visibility=[];offset=0;visible=True
  for line in interface.splitlines(True):
   if re.search(r'=\s*class\b',line):visible=True
   if line.strip() in ['private','strict private','protected','strict protected']:visible=False
   if line.strip() in ['public','published'] or line.strip()=='end;':visible=True
   visibility.append((offset,visible));offset+=len(line)
  for match in re.finditer(r'^\s*(?:(?:class\s+)?(?:function|procedure|constructor|destructor)\s+)',interface,re.M):
   if not next(state for pos,state in reversed(visibility) if pos<=match.start()):continue
   pos=match.start();depth=0;quoted=False;end=pos
   while end<len(interface):
    c=interface[end]
    if c=="'":quoted=not quoted
    if not quoted:
     if c=='(':depth+=1
     elif c==')':depth-=1
     elif c==';' and depth==0:break
    end+=1
   signature=re.sub(r'\s+',' ',interface[pos:end+1]).strip()
   declarations.append(signature)
  if declarations:blocks.append(('code',declarations))
  else:blocks.append(('p','This unit exposes shared data types/constants; inspect its declarations for the exact fields and current version values.'))
 return blocks

def schema_blocks():
 blocks=[];n=2
 purposes={'001':'Foundation project/configuration identity, immutable captures, source facts and authored history.','002':'Additional relational and lifecycle protection beyond base table creation.','003':'Managed cache, retention, release and continuity lifecycle support.','004':'Generation-scoped search indexes and their maintenance.','005':'Document section, audience and related knowledge organization.','006':'Exact revision-link targets and additional integrity constraints.','007':'Attributed author-question metadata and answer provenance.','008':'Revision/locale translation caching and named review support.','009':'Typed authored metadata stored with its exact knowledge revision.'}
 for path in sorted((ROOT/'migrations').glob('*.sql')):
  text=path.read_text(encoding='utf-8-sig');blocks.extend([('section',f'26.{n}. Migration {path.stem.replace("_"," ")}'),('p',purposes[path.name[:3]]+' Source: migrations/'+path.name+'.')]);n+=1
  # Include exact DDL, retaining constraints. Full trigger implementations remain in migrations.
  clean=re.sub(r'--[^\n]*','',text)
  tables=re.findall(r'CREATE\s+(?:VIRTUAL\s+)?TABLE[\s\S]*?;',clean,re.I)
  for table in tables:blocks.append(('code',table.strip().splitlines()))
  alterations=re.findall(r'ALTER\s+TABLE[\s\S]*?;',clean,re.I)
  for statement in alterations:blocks.append(('code',statement.strip().splitlines()))
  indexes=re.findall(r'CREATE\s+(?:UNIQUE\s+)?INDEX\s+(?:IF\s+NOT\s+EXISTS\s+)?([\w]+)',clean,re.I)
  triggers=re.findall(r'CREATE\s+TRIGGER\s+(?:IF\s+NOT\s+EXISTS\s+)?([\w]+)',clean,re.I)
  if indexes:blocks.append(('p','Indexes declared by this migration: '+', '.join(indexes)+'. Inspect their column/order clauses in the named migration when changing query behavior.'))
  if triggers:blocks.append(('p','Triggers declared by this migration: '+', '.join(triggers)+'. Their complete WHEN and RAISE/action bodies are authoritative in the migration; preserve them during migration and backup validation.'))
  if not tables and not alterations:blocks.append(('p','This migration primarily extends constraints/indexes or lifecycle protection rather than introducing a new base table. Review its complete statements before changing related writes.'))
 return blocks

def expand(blocks):
 result=[]
 for kind,value in blocks:
  if kind=='units':result.extend(units_blocks())
  elif kind=='schema':result.extend(schema_blocks())
  else:result.append((kind,value))
 return result

def field(p,kind,instruction=None):
 run=p.add_run();node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),kind);run._r.append(node)
 if instruction:
  node=OxmlElement('w:instrText');node.set(qn('xml:space'),'preserve');node.text=instruction;run._r.append(node)

def document(kind,blocks):
 doc=Document();section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
 section.top_margin=Inches(.72);section.bottom_margin=Inches(.72);section.left_margin=section.right_margin=Inches(.8)
 section.header_distance=section.footer_distance=Inches(.3)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','TOC 1','TOC 2']:
  if name not in doc.styles:doc.styles.add_style(name,1)
  style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
 normal=doc.styles['Normal'];normal.font.size=Pt(10.5 if kind=='User' else 11);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.1;normal.paragraph_format.widow_control=True
 for name,size in [('Title',30),('Subtitle',15),('Heading 1',23),('Heading 2',18),('Heading 3',13),('Caption',10),('TOC 1',11),('TOC 2',10.5)]:doc.styles[name].font.size=Pt(size)
 for name in ['Heading 1','Heading 2','Heading 3']:
  style=doc.styles[name];style.paragraph_format.keep_with_next=True;style.paragraph_format.space_before=Pt(13);style.paragraph_format.space_after=Pt(7)
 doc.styles['Heading 2'].font.color.rgb=RGBColor.from_string('146CB4')
 doc.styles['Heading 3'].font.bold=True
 for name in ['TOC 1','TOC 2']:
  doc.styles[name].paragraph_format.space_after=Pt(2);doc.styles[name].paragraph_format.line_spacing=1.0
 doc.styles['TOC 2'].paragraph_format.left_indent=Inches(.18)
 for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 title='Delphi Project Handoff '+kind+' Guide'
 props=doc.core_properties;props.author='';props.last_modified_by='';props.title=title;props.subject='Detailed application operation' if kind=='User' else 'Architecture contracts implementation and maintenance';props.comments='';props.keywords=''
 doc.add_paragraph(title,'Title');doc.add_paragraph('Explore an unfamiliar Delphi application' if kind=='User' else 'Architecture, data contracts, validation and maintenance','Subtitle');doc.add_paragraph('Experimental release 0.8.0.3  |  Guide edition '+str(VERSION));doc.add_paragraph('Last changed '+CHANGED)
 intro='Open a project, follow its files into screens and actions, and assess what the evidence can show' if kind=='User' else 'From architecture and data ownership to source analysis, storage, testing and release'
 doc.add_paragraph(intro);doc.add_paragraph('Experimental software');doc.add_paragraph('This manual describes Delphi Project Handoff. Information about a selected Delphi application is presented separately by the program.')
 doc.add_page_break();doc.add_paragraph('Contents','Title');field(doc.add_paragraph(),'begin',' TOC \\o "1-2" \\h \\z \\u ');field(doc.add_paragraph(),'separate')
 for k,v in blocks:
  if k in ['part','chapter']:doc.add_paragraph(v,'TOC 1' if k=='part' else 'TOC 2')
 field(doc.add_paragraph(),'end')
 part_pending=False
 for k,v in blocks:
  if k=='part':
   p=doc.add_heading(v,1);p.paragraph_format.page_break_before=v.startswith('PART 1');p.paragraph_format.space_before=Pt(24);part_pending=True
  elif k=='chapter':
   p=doc.add_heading(v,2);p.paragraph_format.page_break_before=(kind=='User' and v.startswith('7 '));p.paragraph_format.space_before=Pt(22) if not part_pending else Pt(13);part_pending=False
  elif k=='section':doc.add_heading(v,3)
  elif k=='p':doc.add_paragraph(v)
  elif k=='steps':
   for i,s in enumerate(v,1):
    p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.22);p.paragraph_format.first_line_indent=Inches(-.22);p.paragraph_format.keep_together=True;p.add_run(str(i)+'. ').bold=True;p.add_run(s)
  elif k=='code':
   for line in v:
    for wrapped in textwrap.wrap(line.expandtabs(2),width=94,replace_whitespace=False,drop_whitespace=False) or [' ']:
     p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.03
     run=p.add_run(wrapped);run.font.name='Consolas';run.font.size=Pt(9)
  elif k=='image':
   source,caption=v.split('|',1);p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True
   width=5.8 if source.startswith('samples/') else 6.8
   picture=p.add_run().add_picture(str(ROOT/source),width=Inches(width));picture._inline.docPr.set('descr',caption)
   doc.add_paragraph(caption,'Caption')
 section.different_first_page_header_footer=True
 header=section.header.paragraphs[0];header.text='Delphi Project Handoff  |  '+kind+' Guide';header.runs[0].font.size=Pt(9)
 footer=section.footer.paragraphs[0];footer.add_run('Guide '+str(VERSION)+'  |  '+CHANGED+'                                    Page ');field(footer,'begin',' PAGE ');field(footer,'separate');footer.add_run('1');field(footer,'end')
 setting=OxmlElement('w:updateFields');setting.set(qn('w:val'),'true');doc.settings.element.append(setting)
 doc.save(ROOT/'docs/guides'/('DPH_'+kind+'_Guide.docx'))

def help_page(blocks):
 parts=['<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Delphi Project Handoff Detailed User Help</title><style>body{font:18px/1.65 Segoe UI,Arial,sans-serif;max-width:1060px;margin:auto;padding:24px;color:#20364a;overflow-wrap:anywhere}h1,h2,h3,h4,a{color:#146cb4}header{border-top:5px solid #f28c28}img{max-width:100%;height:auto}li{margin-bottom:12px}section{margin:32px 0}nav a{display:block}pre{white-space:pre-wrap}@media print{nav{display:none}}</style></head><body><header><h1><span class="notranslate" translate="no">Delphi Project Handoff</span> Detailed User Help</h1><p>Guide edition '+str(VERSION)+' last changed '+CHANGED+'</p></header><nav>']
 chapters=[v for k,v in blocks if k=='chapter']
 for i,title in enumerate(chapters):parts.append(f'<a href="#chapter_{i}">{html.escape(title)}</a>')
 parts.append('</nav>');index=-1
 pattern=r'DelphiProjectHandoff\.exe|\.dphdb(?:\.views)?|learning\.html|\bDPH\b|Open application|Open saved project|Handoff dossier|Learning library|Project map|Screenshots|Application diagrams'
 def prose(s):return re.sub(pattern,lambda m:'<span class="notranslate" translate="no">'+m[0]+'</span>',html.escape(s))
 for k,v in blocks:
  if k=='part':parts.append('<h2>'+html.escape(v)+'</h2>')
  elif k=='chapter':index+=1;parts.append(f'<h3 id="chapter_{index}">'+html.escape(v)+'</h3>')
  elif k=='section':parts.append('<h4>'+html.escape(v)+'</h4>')
  elif k=='p':parts.append('<p>'+prose(v)+'</p>')
  elif k=='steps':parts.append('<ol>'+''.join('<li>'+prose(s)+'</li>' for s in v)+'</ol>')
  elif k=='image':
   source,caption=v.split('|',1);target=ROOT/'help/assets'/Path(source).name;target.write_bytes((ROOT/source).read_bytes())
   parts.append('<figure><img src="assets/'+target.name+'" alt="'+html.escape(caption)+'"><figcaption>'+html.escape(caption)+'</figcaption></figure>')
 parts.append('<p>Optional translation requires an internet connection. Product and code names retain their original spelling.</p><div class="gtranslate_wrapper"></div><script>window.gtranslateSettings={default_language:"en",languages:["en","es","fr","de","it","pt","nl"],wrapper_selector:".gtranslate_wrapper",switcher_horizontal_position:"right",switcher_vertical_position:"top"};</script><script src="https://cdn.gtranslate.net/widgets/latest/float.js" defer></script><script src="assets/translation-selector.js" defer></script></body></html>')
 (ROOT/'help/index.html').write_text('\n'.join(parts),encoding='utf-8')

if __name__=='__main__':
 reports={}
 for kind in ['User','Engineering']:
  blocks=expand(source_blocks('DPH_'+kind+'_Guide.txt'));document(kind,blocks)
  if kind=='User':help_page(blocks)
  reports[kind]={'parts':sum(k=='part' for k,v in blocks),'chapters':sum(k=='chapter' for k,v in blocks),'subsections':sum(k=='section' for k,v in blocks),'words':len(' '.join(' '.join(v) if isinstance(v,list) else v for k,v in blocks).split()),'figures':sum(k=='image' for k,v in blocks)}
 sources={p.relative_to(ROOT).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in [*ROOT.glob('DPH.*.pas'),*ROOT.glob('DPH.UI.*.fmx'),* (ROOT/'migrations').glob('*.sql'),*(ROOT/'docs/guides/source').glob('*.txt'),ROOT/'DelphiProjectHandoff.dproj']}
 (ROOT/'docs/validation/reports/guide-source-validation.json').write_text(json.dumps({'guide_version':VERSION,'last_changed':'2026-10-16','sources':sources,'content':reports,'basis':'Current source and serialized controls, current handoff procedures, exact unit declarations and SQL migration reference'},indent=2)+'\n',encoding='utf-8')
 print(json.dumps(reports,indent=2))
